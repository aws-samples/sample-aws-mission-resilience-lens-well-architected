
import json
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Prompt user for JSON filename ──────────────────────────────────────────────────────────
json_filename = input(
    "Enter the MRL JSON filename (e.g., mrl-v0.5.json): ").strip()

# Extract version from filename
version_match = re.search(r'v(\d+\.\d+)', json_filename)
version = version_match.group(1) if version_match else "unknown"

# ── Load JSON ───────────────────────────────────────────────────────────────────────────────
with open(json_filename, encoding="utf-8") as f:
    data = json.load(f)

wb = Workbook()
ws = wb.active
ws.title = "MRL Assessment"

# ── Color palette ───────────────────────────────────────────────────────────────────────────
COLOR_HEADER_BG = "2E75B6"
COLOR_HEADER_FG = "FFFFFF"
COLOR_PILLAR_BG = "1F4E79"
COLOR_PILLAR_FG = "FFFFFF"
COLOR_QUESTION_BG = "D6E4F0"
COLOR_QUESTION_FG = "1F4E79"
COLOR_ANSWER_BG = "FFFFFF"
COLOR_ALT_BG = "F2F7FB"
COLOR_INPUT_BG = "FFFDE7"
COLOR_SEP_BG = "E8F0F7"

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def make_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def style_cell(cell, bold=False, fg="000000", bg=None, wrap=True,
               size=11, align="left", valign="top"):
    cell.font = Font(name="Calibri", size=size, bold=bold, color=fg)
    if bg:
        cell.fill = make_fill(bg)
    cell.alignment = Alignment(
        wrap_text=wrap, horizontal=align, vertical=valign)
    cell.border = border

# ── Column definitions ──────────────────────────────────────────────────────────────────────
# A  Pillar
# B  Q#
# C  Question
# D  Answer Choice
# E  Improvement Plan
# F  Selected? (customer input)
# G  Notes (customer input)


headers = [
    "Pillar",
    "Q#",
    "Question",
    "Answer Choice",
    "Improvement Plan",
    "Selected? (Y / N / Partial)",
    "Notes / Comments"
]

col_widths = [28, 6, 52, 42, 52, 22, 36]

# ── Write column header row ─────────────────────────────────────────────────────────────────
ws.append(headers)
for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
    cell = ws.cell(row=1, column=col_idx)
    style_cell(cell, bold=True, fg=COLOR_HEADER_FG, bg=COLOR_HEADER_BG,
               size=11, align="center", valign="center")
    ws.column_dimensions[get_column_letter(col_idx)].width = width

ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"

# ── Write data ──────────────────────────────────────────────────────────────────────────────
current_row = 2
q_global = 0

for pillar in data.get("pillars", []):
    pillar_name = pillar["name"]

    for q in pillar.get("questions", []):
        q_global += 1
        q_title = q["title"]
        choices = q.get("choices", [])
        num_choices = len(choices)

        if num_choices == 0:
            continue

        # Track the first row of this question block for merging
        q_start_row = current_row

        for c_idx, c in enumerate(choices):
            choice_title = c.get("title", "")
            ip = c.get("improvementPlan", {})
            ip_text = ip.get("displayText", "") if isinstance(
                ip, dict) else str(ip)

            row_data = [
                # A: Pillar (first row only)
                pillar_name if c_idx == 0 else "",
                f"Q{q_global}" if c_idx == 0 else "",  # B: Q# (first row only)
                # C: Question (first row only)
                q_title if c_idx == 0 else "",
                choice_title,                          # D: Answer Choice
                ip_text,                               # E: Improvement Plan
                "",                                    # F: Selected? – customer fills in
                ""                                     # G: Notes – customer fills in
            ]
            ws.append(row_data)

            bg = COLOR_ALT_BG if c_idx % 2 == 1 else COLOR_ANSWER_BG

            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=current_row, column=col_idx)
                if col_idx == 6:
                    style_cell(cell, bg=COLOR_INPUT_BG,
                               align="center", valign="center")
                elif col_idx == 7:
                    style_cell(cell, bg=COLOR_INPUT_BG,
                               align="left", valign="top")
                elif col_idx in (1, 2, 3):
                    # Pillar, Q#, Question columns – style for merge
                    if c_idx == 0:
                        is_pillar_col = (col_idx == 1)
                        style_cell(
                            cell,
                            bold=True,
                            fg=COLOR_PILLAR_FG if is_pillar_col else COLOR_QUESTION_FG,
                            bg=COLOR_PILLAR_BG if is_pillar_col else COLOR_QUESTION_BG,
                            align="center" if col_idx == 2 else "left",
                            valign="top"
                        )
                    else:
                        # Cells below the first in merged range – style to match
                        is_pillar_col = (col_idx == 1)
                        style_cell(
                            cell,
                            bold=True,
                            fg=COLOR_PILLAR_FG if is_pillar_col else COLOR_QUESTION_FG,
                            bg=COLOR_PILLAR_BG if is_pillar_col else COLOR_QUESTION_BG,
                            align="center" if col_idx == 2 else "left",
                            valign="top"
                        )
                else:
                    style_cell(cell, bg=bg)

            ws.row_dimensions[current_row].height = 40
            current_row += 1

        # Merge Pillar, Q#, and Question columns across all rows for this question
        q_end_row = current_row - 1
        if num_choices > 1:
            ws.merge_cells(
                start_row=q_start_row, start_column=1,
                end_row=q_end_row, end_column=1
            )
            ws.merge_cells(
                start_row=q_start_row, start_column=2,
                end_row=q_end_row, end_column=2
            )
            ws.merge_cells(
                start_row=q_start_row, start_column=3,
                end_row=q_end_row, end_column=3
            )
            # Re-apply alignment to merged cells (openpyxl requires this)
            for col_idx in (1, 2, 3):
                cell = ws.cell(row=q_start_row, column=col_idx)
                is_pillar_col = (col_idx == 1)
                cell.alignment = Alignment(
                    wrap_text=True,
                    horizontal="center" if col_idx == 2 else "left",
                    vertical="top"
                )

        # Thin separator row between questions
        ws.append([""] * 7)
        for col_idx in range(1, 8):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = make_fill(COLOR_SEP_BG)
            cell.border = border
        ws.row_dimensions[current_row].height = 6
        current_row += 1

# ── Instructions tab ────────────────────────────────────────────────────────────────────────
ws_inst = wb.create_sheet("Instructions")
instructions = [
    ("AWS Mission Resilience Lens - Customer Self-Assessment", True, 14),
    ("", False, 11),
    ("How to use this spreadsheet:", True, 12),
    ("", False, 11),
    ("1. Review each question (Column C) and its answer choices (Column D).", False, 11),
    ("2. For each answer choice that applies to your workload, enter Y, N, or Partial in the 'Selected?' column (Column F).", False, 11),
    ("3. Use the Notes column (Column G) to add context, caveats, or action items.", False, 11),
    ("4. Multiple answer choices may apply to a single question.", False, 11),
    ("5. Review the Improvement Plan (Column E) for guidance on each choice.", False, 11),
    ("", False, 11),
    ("Column Guide:", True, 12),
    ("  Pillar            - The MRL pillar this question belongs to.", False, 11),
    ("  Q#                - Question number (Q1-Q39).", False, 11),
    ("  Question          - The full question text.", False, 11),
    ("  Answer Choice     - A specific best practice or answer option.", False, 11),
    ("  Improvement Plan  - AWS guidance for this answer choice.", False, 11),
    ("  Selected?         - Enter Y (yes), N (no), or Partial.", False, 11),
    ("  Notes             - Free-text field for your comments.", False, 11),
    ("", False, 11),
    ("This assessment covers 4 pillars and 39 questions.", False, 11),
    ("AWS recommends conducting this review annually or after significant architectural changes.", False, 11),
]

ws_inst.column_dimensions["A"].width = 90
for text, bold, size in instructions:
    ws_inst.append([text])
    cell = ws_inst.cell(row=ws_inst.max_row, column=1)
    cell.font = Font(name="Calibri", size=size, bold=bold)
    cell.alignment = Alignment(wrap_text=True)
    ws_inst.row_dimensions[ws_inst.max_row].height = 20 if text else 8

# ── Save ────────────────────────────────────────────────────────────────────────────────────
output_path = f"MRL_Assessment_v{version}.xlsx"
wb.save(output_path)
print(f"Done! Processed '{json_filename}' and saved to: {output_path}")
